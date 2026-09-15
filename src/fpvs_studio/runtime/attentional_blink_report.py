"""Durable burst recall records, read-only accuracy queries, and explicit Excel export."""

from __future__ import annotations

import csv
import json
import os
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from pydantic import ValidationError

from fpvs_studio.core.attentional_blink_presets import RECALL_QUESTION_PHASES, RECALL_TASK_ID
from fpvs_studio.core.execution import (
    AttentionalBlinkBurstRecord,
    ParticipantMetadata,
    RunExecutionSummary,
    SessionExecutionSummary,
)
from fpvs_studio.core.paths import logs_dir
from fpvs_studio.core.run_spec import AttentionalBlinkStreamRunSpec
from fpvs_studio.core.session_plan import SessionEntry, SessionPlan
from fpvs_studio.core.task_models import TaskResponseRecord
from fpvs_studio.runtime.reporting_lock import project_reporting_lock

ATTENTIONAL_BLINK_BURSTS_FILENAME = "attentional_blink_bursts_v1.csv"
ATTENTIONAL_BLINK_JOURNAL_FILENAME = "attentional_blink_bursts_v1.jsonl"
ATTENTIONAL_BLINK_ACCURACY_FILENAME = "attentional_blink_accuracy.xlsx"
_SOA_HEADER = (
    "SOA (ms)",
    "Trigger codes",
    "Completed bursts",
    "T1 answers",
    "T1 correct",
    "T1 accuracy (%)",
    "T2 answers",
    "T2 correct",
    "T2 accuracy (%)",
)


class AttentionalBlinkDataError(RuntimeError):
    """An AB research journal cannot be safely read."""


class AttentionalBlinkExportError(RuntimeError):
    """An explicitly requested AB workbook could not be saved."""


@dataclass(frozen=True, slots=True)
class AttentionalBlinkSOASummary:
    soa_ms: float
    burst_count: int
    t1_answer_count: int
    t1_correct_count: int
    t1_accuracy_percent: float | None
    t2_answer_count: int
    t2_correct_count: int
    t2_accuracy_percent: float | None
    condition_trigger_codes: tuple[int, ...] = ()


@dataclass(frozen=True, slots=True)
class AttentionalBlinkDataSummary:
    bursts: tuple[AttentionalBlinkBurstRecord, ...]
    conditions: tuple[AttentionalBlinkSOASummary, ...]
    included_session_count: int
    included_burst_count: int
    total_bursts: int
    warnings: tuple[str, ...] = ()


def _record_key(record: AttentionalBlinkBurstRecord) -> tuple[str | int | None, ...]:
    return (
        record.participant_number,
        record.participant_session_number,
        record.session_id,
        record.run_id,
    )


def _ordered_records(
    records: Iterable[AttentionalBlinkBurstRecord],
) -> tuple[AttentionalBlinkBurstRecord, ...]:
    return tuple(
        sorted(
            records,
            key=lambda row: (
                row.participant_number or "",
                row.participant_session_number or 0,
                row.session_started_at.isoformat() if row.session_started_at else "",
                row.session_id,
                row.burst_number,
            ),
        )
    )


def _read_journal(
    path: Path,
) -> tuple[tuple[AttentionalBlinkBurstRecord, ...], tuple[str, ...]]:
    if not path.exists():
        return (), ()
    try:
        lines = path.read_text(encoding="utf-8").splitlines(keepends=True)
    except (OSError, UnicodeError) as exc:
        raise AttentionalBlinkDataError("The attentional blink journal could not be read.") from exc
    latest: dict[tuple[str | int | None, ...], AttentionalBlinkBurstRecord] = {}
    warnings: list[str] = []
    for index, line in enumerate(lines):
        try:
            payload = json.loads(line)
        except json.JSONDecodeError as exc:
            if index == len(lines) - 1 and not line.endswith("\n"):
                warnings.append(
                    "The final journal write was interrupted. Showing the last complete "
                    "checkpoint for each burst; the unfinished write is not included."
                )
                break
            raise AttentionalBlinkDataError(
                f"The attentional blink journal has invalid JSON on line {index + 1}."
            ) from exc
        try:
            record = AttentionalBlinkBurstRecord.model_validate(payload)
        except ValidationError as exc:
            raise AttentionalBlinkDataError(
                f"The attentional blink journal has an invalid record on line {index + 1}."
            ) from exc
        # Completed test streams follow the same reporting rules as research streams,
        # including journals created before test sessions were included in accuracy.
        latest[_record_key(record)] = record.model_copy(
            update={
                "included_in_accuracy": record.stimulus_completed,
            }
        )
    return _ordered_records(latest.values()), tuple(warnings)


def _summarize_soa(
    records: Iterable[AttentionalBlinkBurstRecord],
) -> tuple[AttentionalBlinkSOASummary, ...]:
    groups: dict[float, list[AttentionalBlinkBurstRecord]] = defaultdict(list)
    for row in records:
        groups[row.requested_soa_ms].append(row)
    summaries = []
    for soa, group in sorted(groups.items()):
        included = [row for row in group if row.included_in_accuracy]
        t1 = [row.t1_correct for row in included if row.t1_correct is not None]
        t2 = [row.t2_correct for row in included if row.t2_correct is not None]
        summaries.append(
            AttentionalBlinkSOASummary(
                soa_ms=soa,
                burst_count=len(included),
                t1_answer_count=len(t1),
                t1_correct_count=sum(t1),
                t1_accuracy_percent=100.0 * sum(t1) / len(t1) if t1 else None,
                t2_answer_count=len(t2),
                t2_correct_count=sum(t2),
                t2_accuracy_percent=100.0 * sum(t2) / len(t2) if t2 else None,
                condition_trigger_codes=tuple(sorted({
                    row.condition_trigger_code for row in group
                })),
            )
        )
    return tuple(summaries)


def load_attentional_blink_data(project_root: Path) -> AttentionalBlinkDataSummary:
    """Read latest burst checkpoints without writing or regenerating any project data.

    A completed stream's valid answers remain usable when a later answer or session
    is aborted. Missing responses are absent from their target's denominator; an
    explicit 'unsure' response is an incorrect answer. Test IDs 0/00 are included
    in accuracy and identified in the burst table. Session inclusion flags in the
    general fixation workbook do not govern this dedicated recall dataset.
    """

    bursts, warnings = _read_journal(logs_dir(project_root) / ATTENTIONAL_BLINK_JOURNAL_FILENAME)
    included = [row for row in bursts if row.included_in_accuracy]
    return AttentionalBlinkDataSummary(
        bursts=bursts,
        conditions=_summarize_soa(bursts),
        included_session_count=len(
            {
                (row.participant_number, row.participant_session_number, row.session_id)
                for row in included
            }
        ),
        included_burst_count=len(included),
        total_bursts=len(bursts),
        warnings=warnings,
    )


class AttentionalBlinkSessionRecorder:
    """Checkpoint target context before playback and each answer before continuing."""

    def __init__(
        self,
        project_root: Path,
        session_plan: SessionPlan,
        *,
        participant_number: str,
        participant_session_number: int | None,
        participant_metadata: ParticipantMetadata | None = None,
        pilot_mode: bool = False,
    ) -> None:
        self._project_root = project_root
        self._plan = session_plan
        self._participant_number = participant_number
        self._participant_metadata = (participant_metadata or ParticipantMetadata()).model_copy(
            deep=True,
        )
        self._pilot_mode = pilot_mode
        self._participant_session_number = participant_session_number
        self._session_started_at = datetime.now(timezone.utc)
        self._records: dict[str, AttentionalBlinkBurstRecord] = {}
        self._soa_repetitions: dict[float, int] = defaultdict(int)
        self._entries = {entry.run_id: entry for entry in session_plan.ordered_entries()}

    def start_entry(self, entry: SessionEntry) -> None:
        if not any(module.task_id == RECALL_TASK_ID for module in entry.post_tasks):
            return
        run = entry.run_spec
        timing = run.attentional_blink
        if not isinstance(timing, AttentionalBlinkStreamRunSpec):
            raise ValueError("AB recall requires a compiled native character stream.")
        targets = {}
        for phase in ("t1", "t2"):
            events = [event for event in run.stimulus_sequence if event.phase == phase]
            if len(events) != 1 or events[0].text is None:
                raise ValueError("AB recall requires exactly one target of each kind per burst.")
            targets[phase] = events[0].text
        markers = [event for event in run.trigger_events if event.label == "condition_start"]
        if len(markers) != 1:
            raise ValueError("AB recall requires one condition onset trigger.")
        self._soa_repetitions[timing.requested_soa_ms] += 1
        cumulative = sum(row.completed_stimulus_s for row in self._records.values())
        self._save(
            AttentionalBlinkBurstRecord(
                project_id=run.project_id,
                participant_number=self._participant_number,
                participant_metadata=self._participant_metadata,
                is_pilot_session=self._pilot_mode,
                participant_session_number=self._participant_session_number,
                session_id=self._plan.session_id,
                run_id=entry.run_id,
                condition_id=entry.condition_id,
                condition_name=entry.condition_name,
                session_seed=self._plan.random_seed,
                run_seed=run.random_seed,
                session_started_at=self._session_started_at,
                burst_number=entry.global_order_index + 1,
                soa_repetition=self._soa_repetitions[timing.requested_soa_ms],
                condition_trigger_code=markers[0].code,
                requested_soa_ms=timing.requested_soa_ms,
                achieved_soa_ms=timing.achieved_soa_ms,
                t1_target=targets["t1"],
                t2_target=targets["t2"],
                planned_duration_s=run.display.total_frames / run.display.refresh_hz,
                cumulative_stimulus_s=cumulative,
            )
        )

    def record_response(self, response: TaskResponseRecord) -> None:
        record = self._records.get(response.run_id)
        if record is None or response.task_id != RECALL_TASK_ID:
            return
        phase = RECALL_QUESTION_PHASES.get(response.question_id or "")
        if phase is None:
            return
        value = response.text_value
        if value is None and response.selected_option_ids:
            value = response.selected_option_ids[0]
        valid = response.valid and not response.aborted and not response.timed_out
        updates: dict[str, object] = {
            f"{phase}_response": value,
            f"{phase}_correct": response.correct if valid else None,
            f"{phase}_valid": valid,
            f"{phase}_rt_ms": (
                response.reaction_time_s * 1000.0 if response.reaction_time_s is not None else None
            ),
            f"{phase}_timed_out": response.timed_out,
            f"{phase}_aborted": response.aborted,
        }
        updated = record.model_copy(update=updates)
        self._save(
            updated.model_copy(
                update={
                    "recall_completed": updated.t1_valid and updated.t2_valid,
                }
            )
        )

    def update_run(self, summary: RunExecutionSummary) -> None:
        record = self._records.get(summary.run_id)
        if record is None:
            return
        run = self._entries[summary.run_id].run_spec
        duration = min(summary.completed_frames, run.display.total_frames) / run.display.refresh_hz
        completed = summary.completed_frames >= run.display.total_frames
        cumulative = (
            sum(
                row.completed_stimulus_s
                for row in self._records.values()
                if row.burst_number < record.burst_number
            )
            + duration
        )
        onsets = {
            onset.sequence_index: onset.time_s for onset in summary.attentional_blink_onsets or ()
        }
        targets = {
            event.phase: onsets.get(event.sequence_index)
            for event in run.stimulus_sequence
            if event.phase in {"t1", "t2"}
        }
        first, second = targets.get("t1"), targets.get("t2")
        self._save(
            record.model_copy(
                update={
                    "run_started_at": summary.started_at,
                    "completed_stimulus_s": duration,
                    "cumulative_stimulus_s": cumulative,
                    "stimulus_completed": completed,
                    "run_aborted": summary.aborted,
                    "included_in_accuracy": completed,
                    "observed_soa_ms": (second - first) * 1000.0
                    if first is not None and second is not None
                    else None,
                }
            )
        )

    def _save(self, record: AttentionalBlinkBurstRecord) -> None:
        path = logs_dir(self._project_root) / ATTENTIONAL_BLINK_JOURNAL_FILENAME
        with project_reporting_lock(self._project_root):
            path.parent.mkdir(parents=True, exist_ok=True)
            if path.exists() and path.stat().st_size:
                with path.open("rb") as previous:
                    previous.seek(-1, os.SEEK_END)
                    if previous.read(1) != b"\n":
                        raise AttentionalBlinkDataError(
                            "The attentional blink journal has an unfinished final write. "
                            "Export the recoverable results and repair that final line before "
                            "recording another session."
                        )
            with path.open("a", encoding="utf-8", newline="\n") as handle:
                handle.write(record.model_dump_json() + "\n")
                handle.flush()
                os.fsync(handle.fileno())
        self._records[record.run_id] = record

    def finish(self, summary: SessionExecutionSummary, *, output_dir: Path | None) -> None:
        if not self._records:
            return
        for record in tuple(self._records.values()):
            self._save(
                record.model_copy(
                    update={
                        "session_finalized": True,
                        "session_aborted": summary.aborted,
                    }
                )
            )
        records = _ordered_records(self._records.values())
        with project_reporting_lock(self._project_root):
            journal = logs_dir(self._project_root) / ATTENTIONAL_BLINK_JOURNAL_FILENAME
            all_records, _warnings = _read_journal(journal)
            _write_burst_csv(
                logs_dir(self._project_root) / ATTENTIONAL_BLINK_BURSTS_FILENAME,
                all_records,
            )
        if output_dir is not None:
            _write_burst_files(output_dir, records)
            for record in records:
                _write_burst_files(output_dir / record.run_id, (record,))


def _csv_value(value: object) -> object:
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _burst_export_payload(record: AttentionalBlinkBurstRecord) -> dict[str, Any]:
    payload = record.model_dump(mode="json", exclude={"participant_metadata"})
    payload["is_test_session"] = record.is_test_session
    for key, value in record.participant_metadata.model_dump(mode="json").items():
        payload[f"participant_{key}"] = json.dumps(value) if isinstance(value, list) else value
    return payload


def _burst_export_fields() -> list[str]:
    return [
        *(key for key in AttentionalBlinkBurstRecord.model_fields if key != "participant_metadata"),
        "is_test_session",
        *(f"participant_{key}" for key in ParticipantMetadata.model_fields),
    ]


def _write_burst_csv(path: Path, records: Iterable[AttentionalBlinkBurstRecord]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(
            "w",
            encoding="utf-8",
            newline="",
            dir=path.parent,
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
            writer = csv.DictWriter(
                handle,
                fieldnames=_burst_export_fields(),
            )
            writer.writeheader()
            for record in records:
                writer.writerow(
                    {key: _csv_value(value) for key, value in _burst_export_payload(record).items()}
                )
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _write_burst_files(directory: Path, records: tuple[AttentionalBlinkBurstRecord, ...]) -> None:
    _write_burst_csv(directory / ATTENTIONAL_BLINK_BURSTS_FILENAME, records)
    (directory / "attentional_blink_bursts_v1.json").write_text(
        json.dumps(
            {"schema_version": "1.0", "bursts": [row.model_dump(mode="json") for row in records]},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def _soa_values(summary: AttentionalBlinkSOASummary) -> tuple[object, ...]:
    return (
        summary.soa_ms,
        ", ".join(str(code) for code in summary.condition_trigger_codes),
        summary.burst_count,
        summary.t1_answer_count,
        summary.t1_correct_count,
        summary.t1_accuracy_percent,
        summary.t2_answer_count,
        summary.t2_correct_count,
        summary.t2_accuracy_percent,
    )


def write_attentional_blink_accuracy_xlsx(
    summary: AttentionalBlinkDataSummary,
    output_path: Path,
) -> Path:
    """Export loaded results with separate SOA, participant/session, and burst tables."""
    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    workbook = Workbook()
    by_soa = workbook.active
    by_soa.title = "Accuracy by SOA"
    by_soa.append(_SOA_HEADER)
    for condition in summary.conditions:
        by_soa.append(_soa_values(condition))
    by_participant = workbook.create_sheet("Participant SOA")
    by_participant.append(
        ("Participant", "Session number", "Session ID", "Test session", *_SOA_HEADER,
         "Pilot session", *ParticipantMetadata.model_fields)
    )
    groups: dict[tuple[str | None, int | None, str], list[AttentionalBlinkBurstRecord]] = (
        defaultdict(list)
    )
    for record in summary.bursts:
        identity = (record.participant_number, record.participant_session_number, record.session_id)
        groups[identity].append(record)
    for identity, records in groups.items():
        for condition in _summarize_soa(records):
            metadata = records[0].participant_metadata.model_dump(mode="json")
            by_participant.append((
                *identity, records[0].is_test_session, *_soa_values(condition),
                records[0].is_pilot_session,
                *(json.dumps(value) if isinstance(value, list) else value
                  for value in metadata.values()),
            ))
    bursts = workbook.create_sheet("Bursts")
    fields = _burst_export_fields()
    bursts.append(fields)
    for record in summary.bursts:
        values = _burst_export_payload(record)
        bursts.append([values[field] for field in fields])
    for worksheet in (by_soa, by_participant, bursts):
        _format_sheet(worksheet)
    notes = workbook.create_sheet("Read me")
    for text in (
        "Attentional blink recall accuracy",
        "One burst row = one participant, session visit, and randomized stream occurrence.",
        "burst_number is 1-based chronological order within a session; "
        "soa_repetition is order within that SOA.",
        "Accuracy uses completed streams and each target's valid answers. "
        "IDs 0/00 are included test sessions, marked in Participant SOA and Bursts.",
        "Pilot sessions retain participant IDs and demographics and are marked separately. "
        "They use local testing checks without EEG hardware. Pilot accuracy is included.",
        "A later session or T2 abort does not discard an earlier valid T1 answer. "
        "Unanswered/invalid responses have null correctness and are excluded from "
        "that target's denominator.",
        "An explicit unsure response counts as incorrect. Separate T1/T2 denominators are shown.",
        "T2 accuracy is unconditional recall, not conditional on a correct T1 response.",
        "General participant/fixation workbook inclusion flags do not alter "
        "this dedicated recall report.",
        "cumulative_stimulus_s includes stream time only; participant response screens "
        "and breaks are excluded.",
        "requested_soa_ms and achieved_soa_ms are compiled onset-to-onset timing; "
        "observed_soa_ms requires actual target flip timestamps.",
        "session_finalized=False means the data is a live or interrupted checkpoint, "
        "not a finalized session.",
        "SOA totals are descriptive answer-weighted accuracy. Use Participant SOA "
        "and Bursts to retain repeated-measure structure for analysis.",
        *summary.warnings,
    ):
        notes.append((text,))
    notes.column_dimensions["A"].width = 115
    for row in notes:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[row[0].row].height = 35
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    except OSError as exc:
        raise AttentionalBlinkExportError(
            "The attentional blink workbook could not be saved."
        ) from exc
    return path


def _format_sheet(worksheet: Any) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet:
        for cell in row:
            cell.alignment = Alignment(vertical="center", horizontal="center")
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.row == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="243A4A")
            elif isinstance(cell.value, float):
                cell.number_format = "0.00"
    for index, column in enumerate(worksheet.columns, 1):
        width = max((len(str(cell.value)) for cell in column if cell.value is not None), default=10)
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 42)
