"""AB recall scoring, durable partial answers, session identity, and exports."""

from __future__ import annotations

import csv
import json

import pytest
from openpyxl import load_workbook
from tests.unit.runtime_launcher_helpers import StubEngine

from fpvs_studio.core.compiler import compile_session_plan
from fpvs_studio.core.enums import ExperimentCategory
from fpvs_studio.core.project_service import build_starter_project
from fpvs_studio.core.task_models import (
    TaskOption,
    TaskQuestionKind,
    TaskResponseKind,
    TaskStepResult,
)
from fpvs_studio.engines.base import TaskEngineInput
from fpvs_studio.runtime.attentional_blink_report import (
    ATTENTIONAL_BLINK_BURSTS_FILENAME,
    ATTENTIONAL_BLINK_JOURNAL_FILENAME,
    AttentionalBlinkDataError,
    load_attentional_blink_data,
    write_attentional_blink_accuracy_xlsx,
)
from fpvs_studio.runtime.run_worker import RuntimeWorker
from fpvs_studio.runtime.task_runner import _score_response


def _plan():
    project = build_starter_project(
        "AB recall",
        experiment_category=ExperimentCategory.ATTENTIONAL_BLINK,
    )
    project.settings.session.block_count = 2
    return compile_session_plan(project, refresh_hz=60.0, random_seed=271)


class RecallEngine(StubEngine):
    def __init__(self, *, abort_t2=False, crash_t2=False, abort_stream=False):
        super().__init__({})
        self.targets = {}
        self.abort_t2 = abort_t2
        self.crash_t2 = crash_t2
        self.abort_stream = abort_stream

    def run_condition(self, run_spec, project_root, **kwargs):
        self.run_id = run_spec.run_id
        self.targets = {
            event.phase: event.text
            for event in run_spec.stimulus_sequence
            if event.phase in {"t1", "t2"}
        }
        result = super().run_condition(run_spec, project_root, **kwargs)
        if self.abort_stream:
            return result.model_copy(
                update={
                    "completed_frames": 60,
                    "aborted": True,
                    "abort_reason": "Stream aborted",
                }
            )
        return result

    def render_task_step(self, step, project_root):
        phase = "t1" if step.question_id == "t1-recall" else "t2"
        if phase == "t2":
            # The first response and target identity must already be on disk while T2 is pending.
            snapshot = next(
                row
                for row in load_attentional_blink_data(project_root).bursts
                if row.run_id == self.run_id and not row.session_finalized
            )
            assert snapshot.t1_response == f" {self.targets['t1']} "
            assert snapshot.t1_correct is True
            assert snapshot.t2_response is None
            assert snapshot.stimulus_completed
            assert not snapshot.session_finalized
            if self.crash_t2:
                raise RuntimeError("Simulated participant screen crash")
            if self.abort_t2:
                return TaskEngineInput(aborted=True)
        return TaskEngineInput(
            text_value=f" {self.targets['t1']} " if phase == "t1" else "unsure",
            reaction_time_s=0.25 if phase == "t1" else 0.75,
        )


def _execute(tmp_path, plan, engine, *, mode="compact", participant="0042", visit=1):
    return RuntimeWorker(engine).execute_session(
        tmp_path,
        plan,
        tmp_path / "runs" / f"session{visit}",
        participant_number=participant,
        participant_session_number=visit,
        runtime_options={
            "export_mode": mode, "serial_enabled": False,
            "verify_refresh_rate": participant not in {"0", "00"},
            "verify_graphics_memory": participant not in {"0", "00"},
        },
    )


@pytest.mark.parametrize("mode", ["compact", "full"])
def test_completed_bursts_score_targets_and_export_chronology(tmp_path, mode):
    plan = _plan()
    engine = RecallEngine()
    result = _execute(tmp_path, plan, engine, mode=mode)
    assert not result.aborted
    transitions = engine._captures["transitions"]
    assert len(transitions) == 6
    assert all(screen["continue_key"] == "space" for screen in transitions)
    assert all(screen["countdown_seconds"] is None for screen in transitions)
    assert all(
        screen["continue_prompt"] == "Press space when you're ready to continue."
        for screen in transitions
    )
    report = load_attentional_blink_data(tmp_path)
    assert report.total_bursts == report.included_burst_count == 6
    assert report.included_session_count == 1
    assert [row.burst_number for row in report.bursts] == [1, 2, 3, 4, 5, 6]
    assert [row.cumulative_stimulus_s for row in report.bursts] == [5, 10, 15, 20, 25, 30]
    assert all(row.session_finalized and row.recall_completed for row in report.bursts)
    for condition in report.conditions:
        assert condition.burst_count == condition.t1_answer_count == condition.t2_answer_count == 2
        assert condition.t1_accuracy_percent == 100
        assert condition.t2_accuracy_percent == 0
        assert condition.condition_trigger_codes == (int(condition.soa_ms / 100),)
        rows = [row for row in report.bursts if row.requested_soa_ms == condition.soa_ms]
        assert [row.soa_repetition for row in rows] == [1, 2]
    for entry, row, summary in zip(
        plan.ordered_entries(), report.bursts, result.run_results, strict=True
    ):
        targets = {
            event.phase: event.text
            for event in entry.run_spec.stimulus_sequence
            if event.phase in {"t1", "t2"}
        }
        assert (row.t1_target, row.t2_target) == (targets["t1"], targets["t2"])
        assert row.t1_target != row.t2_target
        assert row.condition_trigger_code == int(row.requested_soa_ms / 100)
        assert row.requested_soa_ms == row.achieved_soa_ms
        assert row.observed_soa_ms is None
        assert row.t1_rt_ms == 250 and row.t2_rt_ms == 750
        assert summary.task_responses[0].correct is True
        assert summary.task_responses[1].correct is False
        assert summary.task_responses[0].text_value == f" {targets['t1']} "
        assert row.t1_response == f" {targets['t1']} "
        assert row.t2_response == "unsure"
    with (tmp_path / "logs" / ATTENTIONAL_BLINK_BURSTS_FILENAME).open(
        newline="", encoding="utf-8"
    ) as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 6
    assert all(
        row["participant_number"] == "0042" and row["participant_session_number"] == "1"
        for row in rows
    )
    if mode == "compact":
        assert not (tmp_path / "runs").exists()
    else:
        directory = tmp_path / "runs" / "session1"
        payload = json.loads((directory / "attentional_blink_bursts_v1.json").read_text())
        assert len(payload["bursts"]) == 6
        assert (
            directory / plan.ordered_entries()[0].run_id / ATTENTIONAL_BLINK_BURSTS_FILENAME
        ).is_file()
    general_summary = (tmp_path / "logs" / "participant_summary.csv").read_text()
    assert "t1_response" not in general_summary and "t2_response" not in general_summary


def test_t2_abort_preserves_t1_with_separate_denominators(tmp_path):
    result = _execute(tmp_path, _plan(), RecallEngine(abort_t2=True))
    assert result.aborted
    report = load_attentional_blink_data(tmp_path)
    row = report.bursts[0]
    assert row.session_aborted and row.session_finalized
    assert row.t1_correct is True and row.t2_correct is None
    assert row.t2_response is None and row.t2_aborted
    assert row.stimulus_completed and not row.recall_completed
    assert report.included_burst_count == 1
    soa = report.conditions[0]
    assert soa.t1_answer_count == 1 and soa.t1_accuracy_percent == 100
    assert soa.t2_answer_count == 0 and soa.t2_accuracy_percent is None


def test_t2_recall_accuracy_is_retained_when_t1_is_wrong(tmp_path):
    class ReverseRecallEngine(RecallEngine):
        def render_task_step(self, step, project_root):
            digit = "0" if step.question_id == "t1-recall" else self.targets["t2"]
            return TaskEngineInput(text_value=digit, reaction_time_s=0.4)

    _execute(tmp_path, _plan(), ReverseRecallEngine())
    report = load_attentional_blink_data(tmp_path)
    for condition in report.conditions:
        assert condition.t1_accuracy_percent == 0
        assert condition.t2_accuracy_percent == 100
        assert condition.t1_answer_count == condition.t2_answer_count == 2


def test_questionnaire_crash_checkpoints_first_answer_before_second_question(tmp_path):
    plan = _plan()
    # Exercise generic multi-question steps, not only the preset's separate steps.
    for entry in plan.ordered_entries():
        module = entry.post_tasks[0]
        module.steps[0].questions.extend(module.steps[1].questions)
        module.steps = module.steps[:1]
    with pytest.raises(RuntimeError, match="screen crash"):
        _execute(tmp_path, plan, RecallEngine(crash_t2=True))
    report = load_attentional_blink_data(tmp_path)
    row = report.bursts[0]
    assert row.t1_correct is True and row.t1_response == f" {row.t1_target} "
    assert row.t2_correct is None and row.t2_response is None
    assert not row.session_finalized
    assert row.requested_soa_ms in {100, 300, 500}
    assert row.participant_number == "0042" and row.participant_session_number == 1
    checkpoint = next((tmp_path / "logs" / ".task-response-checkpoints").glob("*.jsonl"))
    raw_answers = [json.loads(line) for line in checkpoint.read_text().splitlines()]
    assert len(raw_answers) == 1 and raw_answers[0]["question_id"] == "t1-recall"


def test_aborted_stream_is_visible_but_excluded_from_accuracy(tmp_path):
    _execute(tmp_path, _plan(), RecallEngine(abort_stream=True))
    report = load_attentional_blink_data(tmp_path)
    assert report.total_bursts == 1 and report.included_burst_count == 0
    row = report.bursts[0]
    assert row.run_aborted and not row.stimulus_completed
    assert row.completed_stimulus_s == row.cumulative_stimulus_s == 1
    assert row.t1_response is row.t2_response is None


def test_visit_identity_and_read_only_query_and_workbook(tmp_path):
    plan = _plan()
    _execute(tmp_path, plan, RecallEngine(), visit=1)
    _execute(tmp_path, plan, RecallEngine(), visit=2)
    _execute(tmp_path, plan, RecallEngine(), participant="00", visit=1)
    before = {
        str(path): path.read_bytes() for path in (tmp_path / "logs").rglob("*") if path.is_file()
    }
    report = load_attentional_blink_data(tmp_path)
    assert report.total_bursts == 18 and report.included_burst_count == 18
    assert report.included_session_count == 3
    assert sum(row.is_test_session for row in report.bursts) == 6
    path = write_attentional_blink_accuracy_xlsx(report, tmp_path / "recall.xlsx")
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["Accuracy by SOA", "Participant SOA", "Bursts", "Read me"]
    assert workbook["Bursts"].max_row == 19
    assert workbook["Participant SOA"].max_row == 10
    assert workbook["Bursts"].freeze_panes == "A2"
    assert workbook["Accuracy by SOA"]["F2"].value == 100
    assert workbook["Participant SOA"]["D1"].value == "Test session"
    burst_headers = [cell.value for cell in workbook["Bursts"][1]]
    test_column = burst_headers.index("is_test_session") + 1
    assert (
        sum(workbook["Bursts"].cell(row=index, column=test_column).value for index in range(2, 20))
        == 6
    )
    after = {
        str(path): path.read_bytes() for path in (tmp_path / "logs").rglob("*") if path.is_file()
    }
    assert after == before


def test_soa_summaries_export_sorted_unique_recorded_trigger_codes(tmp_path):
    plan = _plan()
    codes = {100: iter((11, 7)), 300: iter((42, 42)), 500: iter((9, 8))}
    for entry in plan.ordered_entries():
        soa = entry.run_spec.attentional_blink.requested_soa_ms
        code = next(codes[soa])
        entry.run_spec.trigger_events = [
            trigger.model_copy(update={"code": code})
            if trigger.label == "condition_start" else trigger
            for trigger in entry.run_spec.trigger_events
        ]
    _execute(tmp_path, plan, RecallEngine())
    report = load_attentional_blink_data(tmp_path)
    assert [row.soa_ms for row in report.conditions] == [100, 300, 500]
    assert [row.condition_trigger_codes for row in report.conditions] == [
        (7, 11), (42,), (8, 9),
    ]
    assert all(row.t1_answer_count == row.t2_answer_count == 2 for row in report.conditions)
    assert all(row.t1_accuracy_percent == 100 for row in report.conditions)
    assert all(row.t2_accuracy_percent == 0 for row in report.conditions)
    path = write_attentional_blink_accuracy_xlsx(report, tmp_path / "coded_recall.xlsx")
    workbook = load_workbook(path, data_only=True)
    for sheet_name in ("Accuracy by SOA", "Participant SOA"):
        header, *values = list(workbook[sheet_name].values)
        rows = [dict(zip(header, row, strict=True)) for row in values]
        assert [row["SOA (ms)"] for row in rows] == [100, 300, 500]
        assert [row["Trigger codes"] for row in rows] == ["7, 11", "42", "8, 9"]
        assert all(row["T1 accuracy (%)"] == 100 for row in rows)
        assert all(row["T2 accuracy (%)"] == 0 for row in rows)
    with (tmp_path / "logs" / ATTENTIONAL_BLINK_BURSTS_FILENAME).open(
        newline="", encoding="utf-8",
    ) as stream:
        rows = list(csv.DictReader(stream))
    assert "condition_trigger_code" in rows[0]
    assert "condition_trigger_codes" not in rows[0]


@pytest.mark.parametrize("participant", ["0", "00"])
def test_test_mode_always_records_and_calculates_accuracy(tmp_path, participant):
    _execute(tmp_path, _plan(), RecallEngine(), participant=participant)
    report = load_attentional_blink_data(tmp_path)
    assert report.included_burst_count == 6
    assert all(row.is_test_session and row.included_in_accuracy for row in report.bursts)
    assert all(condition.t1_accuracy_percent == 100 for condition in report.conditions)
    assert all(condition.t2_accuracy_percent == 0 for condition in report.conditions)
    with (tmp_path / "logs" / ATTENTIONAL_BLINK_BURSTS_FILENAME).open(
        newline="",
        encoding="utf-8",
    ) as stream:
        assert all(row["is_test_session"] == "True" for row in csv.DictReader(stream))
    # The original journal's exclusion flag must not hide already recorded test answers.
    journal = tmp_path / "logs" / ATTENTIONAL_BLINK_JOURNAL_FILENAME
    payloads = [json.loads(line) for line in journal.read_text().splitlines()]
    for payload in payloads:
        payload["included_in_accuracy"] = False
    journal.write_text("".join(json.dumps(payload) + "\n" for payload in payloads))
    previous_bytes = journal.read_bytes()
    restored = load_attentional_blink_data(tmp_path)
    assert restored.included_burst_count == 6
    assert restored.conditions[0].t1_accuracy_percent == 100
    assert journal.read_bytes() == previous_bytes


@pytest.mark.parametrize(
    "text,expected",
    [
        ("1", True),
        (" 1 ", True),
        ("12", False),
        ("01", False),
        ("unsure", False),
        ("=1", False),
        ("", None),
        ("   ", None),
    ],
)
def test_typed_scoring_matches_whole_answer_without_altering_raw_text(text, expected):
    step = _plan().ordered_entries()[0].post_tasks[0].steps[0]
    question = step.questions[0].model_copy(update={"correct_text": "1"})
    response = TaskStepResult(response_kind=TaskResponseKind.TEXT, text_value=text)
    correct, score = _score_response(step, response, question=question)
    assert correct is expected
    assert score == (float(expected) if expected is not None else None)
    assert response.text_value == text


def test_saved_choice_recall_records_still_score_and_export(tmp_path):
    plan = _plan()
    for entry in plan.ordered_entries():
        targets = {
            event.phase: event.text
            for event in entry.run_spec.stimulus_sequence
            if event.phase in {"t1", "t2"}
        }
        for step in entry.post_tasks[0].steps:
            question = step.questions[0]
            target = targets["t1" if question.question_id == "t1-recall" else "t2"]
            step.questions = [
                question.model_copy(
                    update={
                        "kind": TaskQuestionKind.SINGLE_CHOICE,
                        "correct_text": None,
                        "options": [
                            TaskOption(option_id=digit, label=digit, correct=digit == target)
                            for digit in "0123456789"
                        ],
                    }
                )
            ]

    class ChoiceRecallEngine(RecallEngine):
        def render_task_step(self, step, project_root):
            phase = "t1" if step.question_id == "t1-recall" else "t2"
            return TaskEngineInput(selected_item_ids=(self.targets[phase],))

    _execute(tmp_path, plan, ChoiceRecallEngine())
    report = load_attentional_blink_data(tmp_path)
    assert all(row.t1_correct and row.t2_correct for row in report.bursts)
    assert all(row.t1_response == row.t1_target for row in report.bursts)


def test_missing_journal_does_not_create_files_and_invalid_records_fail_explicitly(tmp_path):
    assert load_attentional_blink_data(tmp_path).total_bursts == 0
    assert not (tmp_path / "logs").exists()
    directory = tmp_path / "logs"
    directory.mkdir()
    (directory / ATTENTIONAL_BLINK_JOURNAL_FILENAME).write_text('{"bad": 4}\n')
    with pytest.raises(AttentionalBlinkDataError, match="invalid record"):
        load_attentional_blink_data(tmp_path)


def test_interrupted_final_write_retains_complete_checkpoint_with_visible_warning(tmp_path):
    _execute(tmp_path, _plan(), RecallEngine(abort_t2=True))
    journal = tmp_path / "logs" / ATTENTIONAL_BLINK_JOURNAL_FILENAME
    with journal.open("a", encoding="utf-8") as stream:
        stream.write('{"schema_version":')
    report = load_attentional_blink_data(tmp_path)
    assert report.bursts[0].t1_correct is True
    assert len(report.warnings) == 1 and "interrupted" in report.warnings[0]
    before = journal.read_bytes()
    with pytest.raises(AttentionalBlinkDataError, match="unfinished final write"):
        _execute(tmp_path, _plan(), RecallEngine(), visit=2)
    assert journal.read_bytes() == before


@pytest.mark.parametrize("mode", ["compact", "full"])
@pytest.mark.parametrize("abort_t2", [False, True])
def test_pilot_demographics_and_answers_survive_exports(tmp_path, mode, abort_t2):
    from fpvs_studio.core.execution import ParticipantMetadata

    metadata = ParticipantMetadata(
        age=24,
        sex="Female",
        handedness="Left handed",
        colorblind=True,
        manual_removed_electrodes=[],
    )
    plan = _plan()
    output = tmp_path / "runs" / "pilot"
    result = RuntimeWorker(RecallEngine(abort_t2=abort_t2)).execute_session(
        tmp_path,
        plan,
        output,
        participant_number="0042",
        participant_session_number=1,
        participant_metadata=metadata,
        runtime_options={
            "pilot_mode": True,
            "export_mode": mode,
            "serial_enabled": False,
            "verify_refresh_rate": False,
            "verify_graphics_memory": False,
        },
    )
    assert result.aborted is abort_t2
    assert result.participant_number == "0042"
    assert result.participant_metadata == metadata
    assert result.runtime_metadata.pilot_mode is True
    report = load_attentional_blink_data(tmp_path)
    assert report.included_burst_count == (1 if abort_t2 else 6)
    assert all(row.is_pilot_session and not row.is_test_session for row in report.bursts)
    assert all(row.participant_metadata == metadata for row in report.bursts)
    assert all(row.t1_correct is True for row in report.bursts)
    assert all(run.runtime_metadata.pilot_mode for run in result.run_results)
    journal = tmp_path / "logs" / ATTENTIONAL_BLINK_JOURNAL_FILENAME
    checkpoints = [json.loads(line) for line in journal.read_text().splitlines()]
    assert all(row["is_pilot_session"] for row in checkpoints)
    assert all(row["participant_metadata"]["age"] == 24 for row in checkpoints)
    with (tmp_path / "logs" / ATTENTIONAL_BLINK_BURSTS_FILENAME).open(newline="") as f:
        rows = list(csv.DictReader(f))
    assert rows[0]["participant_age"] == "24"
    assert rows[0]["participant_sex"] == "Female"
    assert rows[0]["participant_handedness"] == "Left handed"
    assert rows[0]["participant_colorblind"] == "True"
    assert rows[0]["is_pilot_session"] == "True"
    workbook = load_workbook(write_attentional_blink_accuracy_xlsx(report, tmp_path / "pilot.xlsx"))
    for sheet, age_column, pilot_column in [
        ("Bursts", "participant_age", "is_pilot_session"),
        ("Participant SOA", "age", "Pilot session"),
    ]:
        values = list(workbook[sheet].values)
        row = dict(zip(values[0], values[1], strict=True))
        assert row[age_column] == 24
        assert row[pilot_column] is True
    if mode == "compact":
        assert not (tmp_path / "runs").exists()
    else:
        saved = json.loads((output / "session_summary.json").read_text())
        assert saved["participant_metadata"]["colorblind"] is True
        assert saved["runtime_metadata"]["pilot_mode"] is True


def test_journals_before_pilot_mode_remain_readable(tmp_path):
    _execute(tmp_path, _plan(), RecallEngine())
    journal = tmp_path / "logs" / ATTENTIONAL_BLINK_JOURNAL_FILENAME
    records = [json.loads(line) for line in journal.read_text().splitlines()]
    for row in records:
        row.pop("is_pilot_session")
        row.pop("participant_metadata")
    journal.write_text("".join(json.dumps(row) + "\n" for row in records))
    report = load_attentional_blink_data(tmp_path)
    assert report.included_burst_count == 6
    assert all(
        not row.is_pilot_session and row.participant_metadata.is_empty for row in report.bursts
    )
